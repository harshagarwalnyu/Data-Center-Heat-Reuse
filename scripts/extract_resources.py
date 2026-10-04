"""Convert organizer source docs (resources/raw) to plain text (resources/text)."""
from pathlib import Path

import docx
import fitz
import pptx

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "resources" / "raw"
OUT = ROOT / "resources" / "text"


def pdf_text(p: Path) -> str:
    with fitz.open(p) as doc:
        return "\n".join(f"\n--- page {i + 1} ---\n{page.get_text()}" for i, page in enumerate(doc))


def docx_text(p: Path) -> str:
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def pptx_text(p: Path) -> str:
    prs = pptx.Presentation(p)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"\n--- slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            parts.append("[notes] " + slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


HANDLERS = {".pdf": pdf_text, ".docx": docx_text, ".pptx": pptx_text}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for p in sorted(RAW.rglob("*")):
        fn = HANDLERS.get(p.suffix.lower())
        if not fn:
            print(f"SKIP {p.relative_to(RAW)}")
            continue
        text = fn(p)
        dest = OUT / (p.stem.strip().replace(" ", "_") + ".txt")
        dest.write_text(f"SOURCE: {p.relative_to(RAW)}\n\n{text}", encoding="utf-8")
        print(f"{len(text):>9,} chars  {dest.name}")


if __name__ == "__main__":
    main()
